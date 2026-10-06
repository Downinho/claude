"""Gera leads/Leads-Downway.xlsx a partir de leads/leads.json (+ prospects/<slug>/instagram.json, se existir).
Uso: python3 scripts/planilha.py"""
import json, pathlib, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
config = json.load(open(ROOT / 'config.json', encoding='utf-8'))
BASE = config['baseUrl']
leads = json.load(open(ROOT / 'leads' / 'leads.json', encoding='utf-8'))

AREA = {
    'trabalhista': ('direito trabalhista', 'advogado trabalhista em São Paulo'),
    'previdenciario': ('direito previdenciário', 'advogado previdenciário em São Paulo'),
    'familia': ('direito de família', 'advogado de família em São Paulo'),
    'imobiliario': ('direito imobiliário', 'advogado imobiliário em São Paulo'),
    'consumidor': ('direito do consumidor', 'advogado do consumidor em São Paulo'),
    'bancario': ('direito bancário', 'advogado bancário em São Paulo'),
    'dermatologia': ('dermatologia', 'dermatologista em São Paulo'),
    'cardiologia': ('cardiologia', 'cardiologista no Itaim'),
    'endocrinologia': ('endocrinologia', 'endocrinologista em São Paulo'),
    'ortopedia': ('coluna', 'ortopedista de coluna em São Paulo'),
    'nutricao': ('nutrição', 'nutricionista em São Paulo'),
    'nutricao_infantil': ('nutrição materno-infantil', 'nutricionista infantil em São Paulo'),
    'psicologia': ('psicologia', 'psicóloga em São Paulo'),
    'odontologia': ('odontologia', 'dentista no {bairro}'),
    'quiropraxia': ('quiropraxia', 'quiropraxia na zona norte'),
}
REGRA = {'Advogado': 'já seguindo as regras de publicidade da OAB (Provimento 205/2021)',
         'Médico': 'já seguindo a resolução de publicidade médica do CFM (2.336/2023)',
         'Nutricionista': 'já com o CRN em destaque e dentro das regras do conselho',
         'Psicólogo': 'já dentro do código de ética do CFP',
         'Dentista': 'já com CRO e responsável técnico em destaque',
         'Quiropraxia/Fisio': 'com agendamento direto pelo WhatsApp'}
ORG = re.compile(r'^(Advocacia|Clínica|Psicologia|Prez|Odonto|Dentecare|Pandola|SPQuiro|Ninho|Dermic|Dantas)', re.I)


def saudacao(l):
    n = l['nome']
    if re.match(r'^Dra?\.', n):
        return 'Oi, ' + ' '.join(n.split()[:2])
    if l['tipo'] == 'profissional' and not ORG.match(n):
        return 'Oi, ' + n.split()[0]
    return 'Oi, pessoal da ' + n


def mensagem(l, link):
    s = saudacao(l)
    if l['tipo'] == 'profissional':
        area, busca = AREA[l['especialidade']]
        busca = busca.format(bairro=l.get('bairro') or 'São Paulo')
        org = s.startswith('Oi, pessoal')
        para, ate = (f"a {l['nome']}", 'vocês') if org else ('você', 'você')
        return (f"{s}! Sou o [SEU NOME], da Downway. Acompanho {'o conteúdo de vocês' if org else 'seu conteúdo'} de {area} aqui no Instagram "
                f"e montei uma prévia de site para {para}, {REGRA[l['categoria']]}: {link}\n\n"
                f"Hoje, quem pesquisa \"{busca}\" no Google encontra outros profissionais. "
                f"Com o site, essa pessoa chega até {ate} e já fala no WhatsApp. Posso mostrar como fica no ar?")
    return (f"{s}! Sou o [SEU NOME], da Downway. Curto muito o trabalho de vocês aqui no Instagram. "
            f"Montei uma prévia de site para a {l['nome']} só para vocês verem como fica: {link}\n\n"
            f"Quem procura \"{l['categoria'].lower()} {l.get('bairro') or 'São Paulo'}\" no Google cai nos concorrentes. "
            f"Com o site, cai em vocês e já chama no WhatsApp. Abre no celular e me diz o que achou?")


def followup(l):
    return (f"Oi! Conseguiu ver a prévia do site? Se quiser, ajusto cores, fotos ou textos para ficar com a cara "
            f"{'do seu trabalho' if l['tipo'] == 'profissional' else 'de vocês'}. A prévia fica no ar por 7 dias.")


def score(l):
    s = 3 if l['tipo'] == 'profissional' else 1
    s += {'nao encontrado': 3, 'wix gratuito': 2}.get(l['site'], 0)
    seg = l.get('seguidores')
    s += 1 if seg is None else (2 if 1000 <= seg <= 60000 else 0)
    s += 1 if l.get('registro') or l.get('endereco') else 0
    s += 1 if l['prioridade'] == 'A' else 0
    return min(s, 10)


for l in leads:
    f = ROOT / 'prospects' / l['slug'] / 'instagram.json'
    if f.exists():
        ig = json.load(open(f, encoding='utf-8'))
        l['seguidores'] = ig.get('seguidores') or l['seguidores']
        l['link_bio'] = ig.get('link_bio') or ''
        l['ultimo_post'] = ig.get('ultimo_post') or ''
    l['score'] = score(l)
leads.sort(key=lambda l: (-l['score'], l['prioridade'], l['categoria'], l['nome']))

wb = Workbook()
ws = wb.active
ws.title = 'Leads'
COLS = [('#', 5), ('Prior.', 7), ('Score', 7), ('Tipo', 12), ('Categoria', 17), ('Especialidade / nicho', 20), ('Nome', 34),
        ('Instagram', 28), ('Bairro', 16), ('Seguidores', 11), ('Registro', 22), ('Site atual', 15), ('Link da bio', 24),
        ('Último post', 12), ('Prévia', 22), ('Mensagem 1 (DM)', 60), ('Follow-up (2 dias)', 40), ('Status', 13),
        ('Data contato', 13), ('Próximo passo', 24), ('Observação', 36)]
for i, (c, w) in enumerate(COLS, 1):
    ws.cell(1, i, c)
    ws.column_dimensions[get_column_letter(i)].width = w

AZUL = '0F1B2D'
head_fill = PatternFill('solid', fgColor=AZUL)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = head_fill
    c.alignment = Alignment(vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 30

LABEL_ESP = {'nutricao_infantil': 'nutrição infantil', 'previdenciario': 'previdenciário', 'familia': 'família',
             'imobiliario': 'imobiliário', 'bancario': 'bancário', 'nutricao': 'nutrição', 'estetica_auto': 'estética automotiva'}
link_font = Font(color='1F5FBF', underline='single')
thin = Border(bottom=Side(style='thin', color='E6E6E6'))
for r, l in enumerate(leads, 2):
    link = f"{BASE}/{l['slug']}/"
    ig = l['instagram'].lstrip('@')
    esp = l.get('especialidade') or l.get('nicho')
    row = [r - 1, l['prioridade'], l['score'], 'Profissional' if l['tipo'] == 'profissional' else 'Negócio local', l['categoria'],
           LABEL_ESP.get(esp, esp.replace('_', ' ')), l['nome'], '@' + ig, l.get('bairro', ''), l.get('seguidores'),
           l.get('registro', ''), l['site'], l.get('link_bio', ''), l.get('ultimo_post', ''), 'Abrir prévia',
           mensagem(l, link), followup(l), 'Novo', None, '', l.get('observacao', '')]
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        cell.alignment = Alignment(vertical='top', wrap_text=c in (16, 17, 21))
        cell.border = thin
    ws.cell(r, 8).hyperlink = f'https://instagram.com/{ig}'
    ws.cell(r, 8).font = link_font
    ws.cell(r, 15).hyperlink = link
    ws.cell(r, 15).font = link_font
    ws.cell(r, 19).number_format = 'DD/MM/YYYY'
    ws.row_dimensions[r].height = 96

n = len(leads) + 1
ws.freeze_panes = 'H2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{n}'
dv = DataValidation(type='list', formula1='"Novo,Perfil conferido,Contatado,Follow-up enviado,Respondeu,Reunião,Proposta,Fechado,Perdido,Descartado"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'R2:R{n}')
ws.conditional_formatting.add(f'B2:B{n}', CellIsRule(operator='equal', formula=['"A"'], fill=PatternFill('solid', fgColor='D9F2E3'), font=Font(bold=True, color='1E6B3A')))
ws.conditional_formatting.add(f'C2:C{n}', CellIsRule(operator='greaterThanOrEqual', formula=['8'], fill=PatternFill('solid', fgColor='D9F2E3')))
ws.conditional_formatting.add(f'A2:U{n}', FormulaRule(formula=['$R2="Fechado"'], fill=PatternFill('solid', fgColor='C6EFCE')))
ws.conditional_formatting.add(f'A2:U{n}', FormulaRule(formula=['OR($R2="Perdido",$R2="Descartado")'], fill=PatternFill('solid', fgColor='F2F2F2'), font=Font(color='999999')))

# ---------- Resumo (fórmulas ligadas à aba Leads) ----------
rs = wb.create_sheet('Resumo')
rs['A1'] = 'Funil Downway'
rs['A1'].font = Font(bold=True, size=16, color=AZUL)
rs.append([])
rs.append(['Status', 'Qtd'])
for s in ['Novo', 'Perfil conferido', 'Contatado', 'Follow-up enviado', 'Respondeu', 'Reunião', 'Proposta', 'Fechado', 'Perdido', 'Descartado']:
    rs.append([s, f'=COUNTIF(Leads!R:R,"{s}")'])
rs.append(['Total', f'=COUNTA(Leads!G2:G{n})'])
rs.append(['Taxa de resposta', '=IFERROR((B8+B9+B10+B11)/(B6+B7+B8+B9+B10+B11+B12),0)'])
rs['B15'].number_format = '0%'
rs.append(['Faturamento (fechados × preço)', f'=B11*{re.sub(r"[^0-9]", "", config["proposta"]["preco"]) or 0}'])
rs['B16'].number_format = '"R$" #,##0'
rs.append([])
rs.append(['Categoria', 'Leads', 'Prioridade A', 'Fechados'])
cats = sorted({l['categoria'] for l in leads})
for c in cats:
    rs.append([c, f'=COUNTIF(Leads!E:E,"{c}")', f'=COUNTIFS(Leads!E:E,"{c}",Leads!B:B,"A")', f'=COUNTIFS(Leads!E:E,"{c}",Leads!R:R,"Fechado")'])
for row in (3, 18):
    for c in rs[row]:
        if c.value:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head_fill
rs.column_dimensions['A'].width = 32
for col in 'BCD':
    rs.column_dimensions[col].width = 14

# ---------- Como usar ----------
cu = wb.create_sheet('Como usar')
passos = [
    'COMO USAR ESTA PLANILHA',
    '',
    '1. Comece pelos leads de maior Score (já estão no topo). Score vai de 0 a 10.',
    '2. Abra o Instagram do lead (coluna H). Confira: o link da bio NÃO é um site próprio e o último post tem menos de 15 dias. Se não passar, marque "Descartado".',
    '3. Se passou, marque "Perfil conferido" e rode: node generator/instagram.mjs <slug> && npm run build (puxa logo e fotos reais).',
    '4. Suba dist/demo/ para downway.com.br/demo/ e abra a prévia (coluna O) no celular para conferir.',
    '5. Copie a Mensagem 1 (coluna P), troque [SEU NOME] e mande por DM. Marque "Contatado" e a data.',
    '6. Dois dias sem resposta: mande o Follow-up (coluna Q) e marque "Follow-up enviado".',
    '7. Meta: 10 abordagens por dia. Acompanhe a taxa de resposta na aba Resumo.',
    '',
    'REGRAS DE OURO',
    '• Mande as DMs à mão, uma a uma. Ferramenta de disparo em massa faz o Instagram bloquear a conta.',
    '• Advogados, médicos, psicólogos e nutricionistas: o site já sai sem preço, sem depoimento e sem promessa de resultado, com o registro visível. Isso é argumento de venda: "seu site já nasce dentro das regras do conselho".',
    '• Quem disser não ou não responder em 15 dias: apague a prévia do servidor.',
    '• Dados da pesquisa vêm de busca web (06/10/2026). Seguidores e detalhes podem ter mudado: confira no perfil.',
    '',
    'RESPOSTAS PARA OBJEÇÕES',
    '"Já tenho Instagram" → O Instagram fala com quem já te segue. O site é para quem ainda não te conhece e está pesquisando no Google agora.',
    '"Está caro" → Dá menos de R$ 100 por mês no parcelado. Um cliente novo por mês já paga.',
    '"Vou pensar" → Claro! O que faltou para decidir? Se for algo no site, eu ajusto agora.',
    '"Meu conselho não deixa" → O site já segue as regras (registro visível, sem preço, sem promessa, sem depoimento). Posso te mostrar ponto por ponto.',
]
for p in passos:
    cu.append([p])
cu.column_dimensions['A'].width = 150
for r in (1, 11, 17):
    cu.cell(r, 1).font = Font(bold=True, size=13, color=AZUL)
for row in cu.iter_rows():
    row[0].alignment = Alignment(wrap_text=True, vertical='top')

out = ROOT / 'leads' / 'Leads-Downway.xlsx'
wb.save(out)
print(f'{len(leads)} leads -> {out.relative_to(ROOT)} | score>=8: {sum(l["score"] >= 8 for l in leads)}')
